"""Unit tests for resume parser module (PDF, DOCX, TXT) and metadata extractor."""

import io
import docx
import pypdf
import pytest

from src.parsing.resume_parser import (
    ResumeParsingError,
    extract_metadata,
    extract_text_from_docx,
    extract_text_from_txt,
    parse_resume,
)


def test_txt_parsing():
    content = "John Doe\nSoftware Engineer\nPython, Django, AWS\nEmail: john.doe@example.com\nPhone: (555) 123-4567"
    raw_bytes = content.encode("utf-8")

    parsed = parse_resume(raw_bytes, filename="resume.txt")
    assert "John Doe" in parsed["raw_text"]
    assert "python" in parsed["cleaned_text"]


def test_docx_parsing():
    doc = docx.Document()
    doc.add_heading("Jane Smith", level=1)
    doc.add_paragraph("Machine Learning Specialist with PyTorch, NLP, and SQL.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Education"
    table.cell(0, 1).text = "Bachelor of Computer Science"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    parsed = parse_resume(buffer.getvalue(), filename="resume.docx")
    assert "Jane Smith" in parsed["raw_text"]
    assert "Machine Learning" in parsed["raw_text"]
    assert "Education" in parsed["raw_text"]


def test_empty_file_error():
    with pytest.raises(ResumeParsingError):
        parse_resume(b"", filename="empty.txt")


def test_unsupported_format_error():
    with pytest.raises(ResumeParsingError):
        parse_resume(b"dummy data", filename="archive.zip")


def test_extract_metadata():
    text = """
    Alice Johnson
    Email: alice.j@gmail.com
    Phone: +1 415-555-0199
    Education:
    Bachelor of Technology in Information Technology
    Experience:
    5 years of experience in cloud infrastructure and DevOps.
    """
    meta = extract_metadata(text)
    assert "alice.j@gmail.com" in meta["emails"]
    assert any("415" in p for p in meta["phones"])
    assert any("Bachelor" in e for e in meta["education"])
    assert any("5 years" in exp for exp in meta["experience_mentions"])
